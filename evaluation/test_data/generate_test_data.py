"""Generates the synthetic candidate-document fixtures used by the HireLens
evaluation package (``evaluation/test_cases.json`` / ``evaluation/run_evaluation.py``).

All candidates, employers, and job descriptions here are fictional and were
written for this evaluation package. No real person's data is used anywhere
in this file.

This mirrors the pattern already used in ``data/samples/create_sample_cv.py``
(build a DOCX with ``python-docx``) so the evaluation fixtures are generated
the same way the project's own sample CV is, and are fully reproducible by
re-running this script:

    python evaluation/test_data/generate_test_data.py

It also produces the two invalid-input fixtures used by the document-
extraction / error-handling test cases (TC-08): an unsupported file type and
a corrupted DOCX (garbage bytes with a ``.docx`` extension), following the
same construction used in ``tests/test_extraction.py``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUTPUT_DIR = Path(__file__).parent


def _write_cv(
    filename: str,
    *,
    full_name: str,
    contact_line: str,
    summary: str | None,
    skills: str | None,
    experience: list[tuple[str, str, str]],
    education: list[str],
    projects: list[tuple[str, str, str]] | None = None,
    certifications: list[str] | None = None,
) -> None:
    """Build one synthetic candidate CV as a .docx file, matching the
    section structure of data/samples/create_sample_cv.py."""
    document = Document()

    document.add_heading(full_name, 0)
    document.add_paragraph(contact_line)

    if summary:
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(summary)

    if skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(skills)

    if experience:
        document.add_heading("Experience", level=1)
        for role_line, dates, description in experience:
            document.add_heading(role_line, level=2)
            document.add_paragraph(dates)
            document.add_paragraph(description)

    if education:
        document.add_heading("Education", level=1)
        for entry in education:
            document.add_paragraph(entry)

    if projects:
        document.add_heading("Projects", level=1)
        for name, description, technologies in projects:
            document.add_heading(name, level=2)
            document.add_paragraph(description)
            document.add_paragraph(f"Technologies: {technologies}")

    if certifications:
        document.add_heading("Certifications", level=1)
        for cert in certifications:
            document.add_paragraph(cert)

    output_path = OUTPUT_DIR / filename
    document.save(output_path)
    print(f"  wrote {output_path.relative_to(OUTPUT_DIR.parent.parent)}")


def generate_candidate_strong() -> None:
    """TC-01: explicit evidence for nearly every required + preferred item."""
    _write_cv(
        "candidate_strong.docx",
        full_name="Amara Chen",
        contact_line="Remote | amara.chen.eval@example.com",
        summary=(
            "Backend engineer with 6 years of experience building REST APIs "
            "and scalable services."
        ),
        skills="Python, FastAPI, PostgreSQL, Docker, Git, AWS, Kubernetes",
        experience=[
            (
                "Senior Backend Engineer — Nimbus Systems",
                "March 2021 – Present",
                "Built and maintained REST APIs using Python and FastAPI. "
                "Managed PostgreSQL database schemas and query optimization. "
                "Containerized services with Docker and deployed them to "
                "Kubernetes clusters on AWS. Used Git for version control "
                "and code review workflows.",
            ),
            (
                "Backend Developer — Vantage Cloud",
                "June 2019 – February 2021",
                "Developed internal tooling in Python. Used Git for source "
                "control and participated in peer code review.",
            ),
        ],
        education=["BSc Computer Science — State University"],
        projects=[
            (
                "Order Processing API",
                "Built a REST API for order processing and fulfillment.",
                "FastAPI, PostgreSQL, Docker",
            )
        ],
        certifications=["AWS Certified Developer – Associate"],
    )


def generate_candidate_moderate() -> None:
    """TC-02: some major requirements matched, meaningful gaps remain."""
    _write_cv(
        "candidate_moderate.docx",
        full_name="Jordan Ellis",
        contact_line="Austin, TX | jordan.ellis.eval@example.com",
        summary=(
            "Software developer with experience building web applications "
            "and internal tools."
        ),
        skills="Python, Flask, MySQL, Git",
        experience=[
            (
                "Software Developer — Bright Path Retail",
                "August 2020 – Present",
                "Built REST API endpoints using Python and Flask. Queried "
                "and maintained MySQL databases. Used Git for source "
                "control and collaborated in code reviews.",
            )
        ],
        education=["Some college coursework in Information Technology"],
        projects=[
            (
                "Inventory Tracker",
                "An internal tool for tracking warehouse inventory levels.",
                "Flask, MySQL",
            )
        ],
    )


def generate_candidate_low() -> None:
    """TC-03: very limited documented evidence for the target role."""
    _write_cv(
        "candidate_low.docx",
        full_name="Sam Okafor",
        contact_line="Remote | sam.okafor.eval@example.com",
        summary="Recent graduate interested in software development.",
        skills="Java, HTML, CSS",
        experience=[
            (
                "IT Support Intern — Campus Help Desk",
                "Summer 2023",
                "Assisted students and staff with basic computer "
                "troubleshooting and account resets.",
            )
        ],
        education=["Associate Degree in Information Technology — Community College"],
    )


def generate_candidate_sparse() -> None:
    """TC-07: a near-empty CV — just enough text to extract successfully."""
    _write_cv(
        "candidate_sparse.docx",
        full_name="Alex Kim",
        contact_line="alex.kim.eval@example.com",
        summary="Looking for opportunities in tech.",
        skills=None,
        experience=[],
        education=[],
    )


def generate_candidate_variation() -> None:
    """TC-10: differently-structured CV (dense prose, no bullet sections)."""
    document = Document()
    document.add_heading("Riley Morgan", 0)
    document.add_paragraph(
        "Building reliable backend systems for 4+ years. Comfortable across "
        "the stack: Python and Go for services, Postgres and MySQL for "
        "storage, containerizing everything with Docker, and shipping "
        "through CI on AWS. Prior role at Fenwick Data (2020-present) where "
        "day-to-day work included API development, database migrations, "
        "and pair programming; before that, two years at a small agency "
        "doing general web development. Studied Computer Engineering (BEng) "
        "at Northfield Institute of Technology. Comfortable with git-based "
        "workflows and code review culture."
    )
    output_path = OUTPUT_DIR / "candidate_variation.docx"
    document.save(output_path)
    print(f"  wrote {output_path.relative_to(OUTPUT_DIR.parent.parent)}")


def generate_invalid_inputs() -> None:
    """TC-08: an unsupported file type and a corrupted DOCX."""
    txt_path = OUTPUT_DIR / "candidate_unsupported.txt"
    txt_path.write_text(
        "Sam Okafor - Backend Developer\n"
        "This file is intentionally saved as .txt, which HireLens does not "
        "support, to exercise the unsupported-file-type error path."
    )
    print(f"  wrote {txt_path.relative_to(OUTPUT_DIR.parent.parent)}")

    corrupted_path = OUTPUT_DIR / "candidate_corrupted.docx"
    corrupted_path.write_bytes(
        b"This is not a real DOCX file. It has a .docx extension but is "
        b"plain text, to exercise the corrupted-document error path."
    )
    print(f"  wrote {corrupted_path.relative_to(OUTPUT_DIR.parent.parent)}")


def main() -> None:
    print("Generating HireLens evaluation test-data fixtures...")
    generate_candidate_strong()
    generate_candidate_moderate()
    generate_candidate_low()
    generate_candidate_sparse()
    generate_candidate_variation()
    generate_invalid_inputs()
    print("Done.")


if __name__ == "__main__":
    main()
