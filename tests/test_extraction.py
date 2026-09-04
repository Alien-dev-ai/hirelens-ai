"""Tests for src/extraction/document_extractor.py."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document

from src.extraction.document_extractor import extract_document

# --- PDF test helper -------------------------------------------------------
#
# We need a valid, minimal PDF to exercise the real PDF-extraction path
# without pulling in a heavyweight PDF-generation dependency. pypdf (already
# a project dependency, used for *reading*) has no simple text-writing API,
# so instead we hand-write the PDF bytes directly: PDF is a plain-text-ish
# format, and a single-page document with a text content stream is only a
# few dozen lines of the spec's own syntax. This keeps the test suite free
# of any additional dependency.


def _make_minimal_pdf_bytes(text: str) -> bytes:
    """Build the raw bytes of a minimal, valid single-page PDF containing `text`."""
    lines = text.split("\n")
    content_lines = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    for i, line in enumerate(lines):
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        if i == 0:
            content_lines.append(f"({escaped}) Tj")
        else:
            content_lines.append("T*")
            content_lines.append(f"({escaped}) Tj")
    content_lines.append("ET")
    content = "\n".join(content_lines).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
    ]

    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(buf.tell())
        buf.write(f"{i} 0 obj\n".encode("latin-1"))
        buf.write(obj)
        buf.write(b"\nendobj\n")
    xref_offset = buf.tell()
    n = len(objects) + 1
    buf.write(f"xref\n0 {n}\n".encode("latin-1"))
    buf.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buf.write(f"{off:010d} 00000 n \n".encode("latin-1"))
    buf.write(b"trailer\n")
    buf.write(f"<< /Size {n} /Root 1 0 R >>\n".encode("latin-1"))
    buf.write(b"startxref\n")
    buf.write(f"{xref_offset}\n".encode("latin-1"))
    buf.write(b"%%EOF")
    return buf.getvalue()


def _make_blank_pdf_bytes() -> bytes:
    """Build a minimal, valid single-page PDF with no text content at all."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << >> /MediaBox [0 0 612 792] >>",
    ]
    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(buf.tell())
        buf.write(f"{i} 0 obj\n".encode("latin-1"))
        buf.write(obj)
        buf.write(b"\nendobj\n")
    xref_offset = buf.tell()
    n = len(objects) + 1
    buf.write(f"xref\n0 {n}\n".encode("latin-1"))
    buf.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buf.write(f"{off:010d} 00000 n \n".encode("latin-1"))
    buf.write(b"trailer\n")
    buf.write(f"<< /Size {n} /Root 1 0 R >>\n".encode("latin-1"))
    buf.write(b"startxref\n")
    buf.write(f"{xref_offset}\n".encode("latin-1"))
    buf.write(b"%%EOF")
    return buf.getvalue()


# --- Tests -------------------------------------------------------------


def test_extract_valid_docx(tmp_path: Path):
    """A well-formed DOCX with real paragraph text should extract successfully."""
    docx_path = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Jordan Rivera")
    document.add_paragraph("Backend Engineer with 5 years of Python experience.")
    document.add_paragraph("")  # empty paragraph should be ignored
    document.save(docx_path)

    result = extract_document(docx_path)

    assert result.extraction_success is True
    assert result.error_message is None
    assert result.filename == "candidate.docx"
    assert result.file_type == "docx"
    assert "Jordan Rivera" in result.text
    assert "Backend Engineer with 5 years of Python experience." in result.text


def test_extract_valid_pdf(tmp_path: Path):
    """A well-formed PDF with real text content should extract successfully."""
    pdf_path = tmp_path / "candidate.pdf"
    pdf_path.write_bytes(
        _make_minimal_pdf_bytes("Jordan Rivera\nBackend Engineer with Python experience")
    )

    result = extract_document(pdf_path)

    assert result.extraction_success is True
    assert result.error_message is None
    assert result.filename == "candidate.pdf"
    assert result.file_type == "pdf"
    assert "Jordan Rivera" in result.text
    assert "Backend Engineer with Python experience" in result.text


def test_unsupported_file_type_returns_failure(tmp_path: Path):
    """An unsupported extension (e.g. .txt) should fail gracefully, not raise."""
    txt_path = tmp_path / "candidate.txt"
    txt_path.write_text("Jordan Rivera, Backend Engineer")

    result = extract_document(txt_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None
    assert "unsupported" in result.error_message.lower()


def test_missing_file_returns_failure(tmp_path: Path):
    """A path that does not exist on disk should fail gracefully, not raise."""
    missing_path = tmp_path / "does_not_exist.pdf"

    result = extract_document(missing_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None
    assert "not found" in result.error_message.lower()


def test_empty_docx_returns_failure(tmp_path: Path):
    """A DOCX with only whitespace/empty paragraphs should fail as 'no text'."""
    docx_path = tmp_path / "empty.docx"
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.save(docx_path)

    result = extract_document(docx_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None


def test_blank_pdf_returns_failure(tmp_path: Path):
    """A structurally valid PDF with no text content should fail as 'no text'."""
    pdf_path = tmp_path / "blank.pdf"
    pdf_path.write_bytes(_make_blank_pdf_bytes())

    result = extract_document(pdf_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None


def test_malformed_pdf_does_not_crash(tmp_path: Path):
    """A file with a .pdf extension but garbage content must fail gracefully."""
    bad_pdf_path = tmp_path / "corrupted.pdf"
    bad_pdf_path.write_bytes(b"this is not a real PDF file at all")

    result = extract_document(bad_pdf_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None


def test_malformed_docx_does_not_crash(tmp_path: Path):
    """A file with a .docx extension but garbage (non-zip) content must fail gracefully."""
    bad_docx_path = tmp_path / "corrupted.docx"
    bad_docx_path.write_bytes(b"this is not a real DOCX file at all")

    result = extract_document(bad_docx_path)

    assert result.extraction_success is False
    assert result.text == ""
    assert result.error_message is not None


def test_extract_document_accepts_string_path(tmp_path: Path):
    """extract_document should accept a plain string path, not just a Path object."""
    docx_path = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Sam Lee")
    document.save(docx_path)

    result = extract_document(str(docx_path))

    assert result.extraction_success is True
    assert "Sam Lee" in result.text
