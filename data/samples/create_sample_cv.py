from docx import Document

document = Document()

document.add_heading("John Doe", 0)

document.add_paragraph(
    "Lagos, Nigeria | john.doe@example.com | +1 555 123 4567"
)

document.add_heading("Professional Summary", level=1)
document.add_paragraph(
    "Software Engineer with experience building backend applications, "
    "REST APIs, and database-driven systems."
)

document.add_heading("Skills", level=1)
document.add_paragraph(
    "Python, FastAPI, PostgreSQL, Git, Docker"
)

document.add_heading("Experience", level=1)

document.add_heading("Software Engineer — Tech Solutions Ltd", level=2)
document.add_paragraph("January 2022 – Present")

document.add_paragraph(
    "Developed REST APIs using Python and FastAPI.\n"
    "Worked with PostgreSQL databases.\n"
    "Built containerized applications using Docker.\n"
    "Collaborated with engineering teams using Git for version control."
)

document.add_heading("Junior Developer — Startup Labs", level=2)
document.add_paragraph("June 2020 – December 2021")

document.add_paragraph(
    "Assisted with backend application development.\n"
    "Supported software development workflows and version control."
)

document.add_heading("Education", level=1)

document.add_paragraph(
    "BSc Computer Science — University of Lagos"
)

document.add_heading("Projects", level=1)

document.add_heading("Task Management API", level=2)

document.add_paragraph(
    "Built a REST API for managing tasks using FastAPI and PostgreSQL."
)

document.add_paragraph(
    "Technologies: FastAPI, PostgreSQL"
)

output_path = "data/samples/john_doe_cv.docx"

document.save(output_path)

print(f"Sample CV created successfully: {output_path}")