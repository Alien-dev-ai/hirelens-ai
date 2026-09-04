"""Document extraction for HireLens AI.

Converts an uploaded candidate document (PDF or DOCX) into plain text and
returns it as an :class:`~src.models.schemas.ExtractedDocument`.

This module is intentionally independent of any UI framework (e.g.
Streamlit) so it can be reused for both single-file and batch processing,
and it never raises for ordinary user-input problems (missing file,
unsupported type, corrupted document, empty document) — those are all
reported back via ``ExtractedDocument.extraction_success`` and
``error_message`` instead.
"""

from __future__ import annotations

from pathlib import Path

from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from src.models.schemas import ExtractedDocument

# Extensions this module knows how to extract text from.
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_document(file_path: str | Path) -> ExtractedDocument:
    """Extract text from a candidate document.

    Detects the file type from the extension (``.pdf`` or ``.docx``),
    dispatches to the appropriate extractor, and validates the result.

    Args:
        file_path: Path (or path-like string) to the document to extract.

    Returns:
        An ``ExtractedDocument`` describing the outcome. On any failure —
        missing file, unsupported extension, unreadable/corrupted document,
        or a document that yields no meaningful text — ``extraction_success``
        is ``False`` and ``error_message`` explains why. This function does
        not raise for those cases.
    """
    path = Path(file_path)
    filename = path.name
    file_type = path.suffix.lower().lstrip(".")

    if not path.exists() or not path.is_file():
        return ExtractedDocument(
            filename=filename,
            file_type=file_type,
            text="",
            extraction_success=False,
            error_message=f"File not found: {path}",
        )

    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return ExtractedDocument(
            filename=filename,
            file_type=file_type,
            text="",
            extraction_success=False,
            error_message=f"Unsupported file type: '{extension or '(none)'}'. "
            f"Supported types are: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.",
        )

    if extension == ".pdf":
        text, error_message = _extract_pdf(path)
    else:  # extension == ".docx"
        text, error_message = _extract_docx(path)

    if error_message is not None:
        return ExtractedDocument(
            filename=filename,
            file_type=file_type,
            text="",
            extraction_success=False,
            error_message=error_message,
        )

    if not text.strip():
        return ExtractedDocument(
            filename=filename,
            file_type=file_type,
            text="",
            extraction_success=False,
            error_message="No extractable text was found in the document.",
        )

    return ExtractedDocument(
        filename=filename,
        file_type=file_type,
        text=text,
        extraction_success=True,
        error_message=None,
    )


def _extract_pdf(path: Path) -> tuple[str, str | None]:
    """Extract text from all pages of a PDF file.

    Returns:
        A ``(text, error_message)`` tuple. ``error_message`` is ``None`` on
        success (though ``text`` may still be empty if the PDF has no
        extractable text, which is handled by the caller).
    """
    try:
        reader = PdfReader(path)
    except (PdfReadError, OSError, ValueError) as exc:
        return "", f"Failed to read PDF file: {exc}"

    if reader.is_encrypted:
        try:
            # Try an empty password in case the PDF is trivially "encrypted".
            reader.decrypt("")
        except Exception:
            pass
        if reader.is_encrypted:
            return "", "PDF is encrypted/password-protected and could not be read."

    page_texts: list[str] = []
    try:
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                page_texts.append(page_text.strip())
    except Exception as exc:  # pypdf can raise a variety of parsing errors
        return "", f"Failed to extract text from PDF: {exc}"

    return "\n\n".join(page_texts), None


def _extract_docx(path: Path) -> tuple[str, str | None]:
    """Extract text from the paragraphs of a DOCX file, ignoring empty ones.

    Returns:
        A ``(text, error_message)`` tuple. ``error_message`` is ``None`` on
        success (though ``text`` may still be empty if the DOCX has no
        non-empty paragraphs, which is handled by the caller).
    """
    try:
        document = Document(path)
    except (PackageNotFoundError, OSError, ValueError, KeyError) as exc:
        return "", f"Failed to read DOCX file: {exc}"
    except Exception as exc:  # defend against other corrupt-file failures
        return "", f"Failed to read DOCX file: {exc}"

    try:
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    except Exception as exc:
        return "", f"Failed to extract text from DOCX: {exc}"

    return "\n".join(paragraphs), None
